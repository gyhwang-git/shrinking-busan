from __future__ import annotations

import csv
import json
import math
from collections import defaultdict
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
DATA = ROOT / "data"
DERIVED = DATA / "derived"
PREVIEW = ROOT / "preview_202609.html"
DISTRICTS = ["남구", "북구", "사상구", "영도구"]
BUSAN_NAME = "부산 전체"
BUSAN_DISTRICTS = ["중구", "서구", "동구", "영도구", "부산진구", "동래구", "남구", "북구", "해운대구", "사하구", "금정구", "강서구", "연제구", "수영구", "사상구", "기장군"]
PROFILES = DISTRICTS + [BUSAN_NAME]
DISTRICT_CODES = {"남구": "210700", "북구": "210800", "사상구": "211500", "영도구": "210400"}
METRIC_KEYS = ["population_growth", "worker_growth", "living_ratio", "persistence", "connectivity"]


def number(value):
    if value is None:
        return None
    text = str(value).strip().replace(",", "").replace("'", "")
    if not text or text in {"-", "X"}:
        return None
    try:
        return float(text)
    except ValueError:
        return None


def find_file(pattern):
    matches = sorted(DATA.glob(pattern))
    if not matches:
        raise FileNotFoundError(f"Required input not found: {pattern}")
    return matches[0]


def resident_population(year):
    path = find_file(f"101_DT_1B04005N_M_{year}_PIVOT.csv")
    populations = {district: 0.0 for district in PROFILES}
    with path.open("r", encoding="cp949", newline="") as handle:
        reader = csv.reader(handle)
        next(reader, None)  # KOSIS table title
        next(reader, None)  # blank line
        header = next(reader)
        value_index = next(
            (i for i, col in enumerate(header) if col.startswith(f"{year}12_총인구수")),
            None,
        )
        if value_index is None:
            raise ValueError(f"December total population column not found in {path.name}")
        for row in reader:
            if len(row) <= value_index or len(row) < 4:
                continue
            raw_area = row[1].strip()
            district = BUSAN_NAME if raw_area == "부산광역시" else raw_area
            if district not in populations or not row[0].strip().startswith("'26"):
                continue
            if row[2].strip() != "'0" or row[3].strip() != "계":
                continue
            value = number(row[value_index])
            if value is not None:
                populations[district] += value
    if any(value <= 0 for value in populations.values()):
        raise ValueError(f"Some 12/31/{year} district populations were not read: {populations}")
    return populations


def employee_totals(year):
    path = find_file(f"읍면동별 산업대분류별 현황_{year}년기준자료.xlsx")
    ws = load_workbook(path, read_only=True, data_only=True).active
    totals = {}
    for row in ws.iter_rows(min_row=5, max_col=4, values_only=True):
        if not row or row[0] is None or row[1] is None:
            continue
        area = str(row[0]).strip()
        industry = str(row[1]).strip()
        if area == "21.부산광역시" and industry in {"전산업", "전체산업"}:
            totals[BUSAN_NAME] = number(row[3])
        for district, code in DISTRICT_CODES.items():
            if area.startswith(code + ".") and industry in {"전산업", "전체산업"}:
                totals[district] = number(row[3])
    missing = [district for district in PROFILES if totals.get(district) is None]
    if missing:
        raise ValueError(f"Total industry worker counts missing for {year}: {missing}")
    return totals


def living_population_2024():
    path = find_file("*생활인구*월별*일평균.xlsx")
    ws = load_workbook(path, read_only=True, data_only=True).active
    # Sum the three mutually exclusive population categories by hour, district and month.
    hourly = defaultdict(float)
    hour_counts = defaultdict(int)
    month_counts = defaultdict(int)
    for row in ws.iter_rows(min_row=2, max_col=7, values_only=True):
        if not row or row[0] is None or row[2] is None:
            continue
        month = str(row[0]).strip()
        if not month.startswith("2024"):
            continue
        dong_name = str(row[2]).strip()
        district = dong_name.split(" ", 1)[0]
        if district not in BUSAN_DISTRICTS:
            continue
        hour = str(row[3]).strip()
        value = sum(number(row[i]) or 0 for i in (4, 5, 6))
        for profile in ([district] if district in DISTRICTS else []) + [BUSAN_NAME]:
            key = (profile, month, hour)
            hourly[key] += value
            hour_counts[key] += 1
            month_counts[(profile, month)] += 1

    monthly_average = {district: {} for district in PROFILES}
    for district in PROFILES:
        months = sorted(month for d, month in month_counts if d == district)
        if len(months) != 12:
            raise ValueError(f"Expected 12 living population months for {district}; found {months}")
        for month in months:
            hour_values = [hourly.get((district, month, f"{hour:02d}시"), 0) for hour in range(24)]
            observed_hours = sum((district, month, f"{hour:02d}시") in hourly for hour in range(24))
            if observed_hours != 24:
                raise ValueError(f"Expected 24 hourly values for {district}, {month}")
            monthly_average[district][month] = sum(hour_values) / 24
    annual_average = {district: sum(monthly_average[district].values()) / 12 for district in PROFILES}
    active_months = {
        district: sum(value > 0 for value in monthly_average[district].values()) for district in PROFILES
    }
    return annual_average, active_months, monthly_average


def migration_connectivity():
    path = find_file("*인구이동*.csv")
    results = {}
    with path.open("r", encoding="cp949", newline="") as handle:
        reader = csv.reader(handle)
        for _ in range(4):
            next(reader, None)  # grouped headers and total-row labels
        for row in reader:
            if len(row) < 16:
                continue
            profile = BUSAN_NAME if row[0].strip() == "합계" else row[0].strip()
            if profile not in PROFILES:
                continue
            # Annual 2024 resident moves: total inflow/outflow (cols 1, 4),
            # inter-gu inflow/outflow inside Busan (cols 10, 13).
            total_in, total_out = number(row[1]), number(row[4])
            external_in, external_out = number(row[10]), number(row[13])
            if None in (total_in, total_out, external_in, external_out):
                continue
            results[profile] = {
                "total_moves": total_in + total_out,
                "external_moves": external_in + external_out,
                "share": (external_in + external_out) / (total_in + total_out) * 100
                if total_in + total_out else 0,
            }
    if any(district not in results for district in PROFILES):
        raise ValueError(f"2024 inter-gu migration rows missing: {results}")
    return results


def annualized_change(start, end, years=3):
    if start <= 0 or end <= 0:
        return None
    return ((end / start) ** (1 / years) - 1) * 100


def score_values(records, key):
    values = [records[district][key] for district in PROFILES]
    low, high = min(values), max(values)
    if math.isclose(low, high, rel_tol=1e-10, abs_tol=1e-10):
        return {district: 50 for district in PROFILES}
    return {district: round((records[district][key] - low) / (high - low) * 100, 1) for district in PROFILES}


def make_data():
    population_2021 = resident_population(2021)
    population_2024 = resident_population(2024)
    workers_2021 = employee_totals(2021)
    workers_2024 = employee_totals(2024)
    living_avg, active_months, monthly = living_population_2024()
    mobility = migration_connectivity()
    low_season_ratio = {
        district: sum(sorted(monthly[district].values())[:3]) / 3 / living_avg[district] * 100
        for district in PROFILES
    }

    records = {}
    for district in PROFILES:
        records[district] = {
            "population_2021": population_2021[district],
            "population_2024": population_2024[district],
            "population_growth": annualized_change(population_2021[district], population_2024[district]),
            "workers_2021": workers_2021[district],
            "workers_2024": workers_2024[district],
            "worker_growth": annualized_change(workers_2021[district], workers_2024[district]),
            "living_population_2024": living_avg[district],
            "living_ratio": living_avg[district] / population_2024[district] * 100,
            "active_months": active_months[district],
            "persistence": low_season_ratio[district],
            "connectivity": mobility[district]["share"],
            "external_moves": mobility[district]["external_moves"],
            "total_moves": mobility[district]["total_moves"],
            "monthly_living_population": monthly[district],
        }
    for key in METRIC_KEYS:
        scores = score_values(records, key)
        for district in PROFILES:
            records[district].setdefault("scores", {})[key] = scores[district]
    return {
        "period": "2015/2024 DID change; population and worker change 2021-2024; living population and migration 2024",
        "districts": {district: records[district] for district in DISTRICTS},
        "busan": records[BUSAN_NAME],
        "score_method": "Min-max standardized to 0-100 across four displayed districts and Busan overall; equal raw values receive 50.",
        "definitions": {
            "population_growth": "2021 and 2024 December registered population, annualized over three years.",
            "worker_growth": "2021 and 2024 total-industry workers, annualized over three years.",
            "living_ratio": "Average 2024 monthly 24-hour mean of residential+workplace+visitor population divided by 2024 December registered population.",
            "persistence": "Mean of the three lowest monthly living population averages divided by the 12-month average, multiplied by 100.",
            "connectivity": "Inter-gu inflow and outflow divided by all inflow and outflow in 2024 resident migration data; this is a residential-move proxy, not daily living-movement OD.",
        },
    }


def svg_for(record, color, busan_record):
    cx, cy, radius = 280, 205, 126
    angles = [-90, -18, 54, 126, 198]
    def point(angle, value):
        rad = math.radians(angle)
        r = radius * value / 100
        return (cx + math.cos(rad) * r, cy + math.sin(rad) * r)
    rings = []
    for level in (20, 40, 60, 80, 100):
        pts = [point(angle, level) for angle in angles]
        rings.append(f'<polygon points="{" ".join(f"{x:.1f},{y:.1f}" for x,y in pts)}" class="radar-grid" />')
    axes = []
    for angle in angles:
        x, y = point(angle, 100)
        axes.append(f'<line x1="{cx}" y1="{cy}" x2="{x:.1f}" y2="{y:.1f}" class="radar-axis" />')
    values = [record["scores"][key] for key in METRIC_KEYS]
    pts = [point(angle, value) for angle, value in zip(angles, values)]
    polygon = " ".join(f"{x:.1f},{y:.1f}" for x, y in pts)
    marks = "".join(f'<circle cx="{x:.1f}" cy="{y:.1f}" r="4" class="radar-dot" />' for x, y in pts)
    busan_values = [busan_record["scores"][key] for key in METRIC_KEYS]
    busan_pts = [point(angle, value) for angle, value in zip(angles, busan_values)]
    busan_polygon = " ".join(f"{x:.1f},{y:.1f}" for x, y in busan_pts)
    busan_marks = "".join(f'<circle cx="{x:.1f}" cy="{y:.1f}" r="3.2" class="radar-busan-dot" />' for x, y in busan_pts)
    labels = ["①정주인구 변화", "②고용기반 변화", "③체류인구 흡인력", "④생활인구 지속성", "⑤광역 이동성"]
    label_markup = []
    for angle, label in zip(angles, labels):
        x, y = point(angle, 137)
        anchor = "middle" if abs(math.cos(math.radians(angle))) < .35 else ("start" if math.cos(math.radians(angle)) > 0 else "end")
        label_markup.append(f'<text x="{x:.1f}" y="{y:.1f}" text-anchor="{anchor}" class="radar-label">{label}</text>')
    return (
        '<svg class="radar-svg" viewBox="0 0 560 400" role="img" aria-label="축소양상 오각형 지표 비교">'
        + "".join(rings + axes)
        + f'<polygon points="{polygon}" class="radar-shape" style="--radar-color:{color}" />'
        + marks
        + f'<polygon points="{busan_polygon}" class="radar-busan-shape"><title>부산 전체 기준선</title></polygon>'
        + busan_marks + "".join(label_markup)
        + '</svg>'
    )


def html_section(data):
    palette = {"남구": "#2d8763", "북구": "#397db3", "사상구": "#d17a4d", "영도구": "#9564a8"}
    rows = []
    district_metrics = {}
    for district in DISTRICTS:
        record = data["districts"][district]
        metrics = [
            ("① ’21~’24 연평균 인구증감률", f'{record["population_growth"]:+.2f}%'),
            ("② ’21~’24 연평균 종사자 증감률", f'{record["worker_growth"]:+.2f}%'),
            ("③ ’24 생활인구/주민등록 인구 비율", f'{record["living_ratio"]:.1f}%'),
            ("\\u2463 \\uC0DD\\uD65C\\uC778\\uAD6C \\uC9C0\\uC18D\\uC131", f'{record["persistence"]:.1f}%'),
            ("⑤ 구/군간 전입/전출 비율", f'{record["connectivity"]:.1f}%'),
        ]
        district_metrics[district] = [value for _, value in metrics]
        rows.append(
            f'<article class="shrinkage-radar-card" style="--district-color:{palette[district]}">'
            f'<header><h3>{district}</h3><small>공통 척도 0–100 · 최대 100</small></header>'
            f'{svg_for(record, palette[district], data["busan"])}'
            '</article>'
        )
    table_labels = [
        ("① 정주인구 변화", "’21~’24 연평균 인구증감률"),
        ("② 고용기반 변화", "’21~’24 연평균 종사자 증감률"),
        ("③ 체류인구 흡인력", "’24 생활인구/주민등록 인구 비율"),
        ("④ 생활인구 지속성", "월별 생활인구 하위 3개월 평균 ÷ 연평균 생활인구"),
        ("⑤ 광역 이동성", "구/군간 전입/전출 비율"),
    ]
    table_rows = "".join(
        f'<tr><th>{label}</th><td class="shrinkage-indicator">{indicator}</td>'
        + "".join(f'<td>{district_metrics[district][index]}</td>' for district in DISTRICTS)
        + "</tr>"
        for index, (label, indicator) in enumerate(table_labels)
    )
    table = (
        '<div class="shrinkage-metrics-table-wrap"><table class="shrinkage-metrics-table"><thead><tr><th>진단 항목</th><th>대표 분석 지표</th>'
        + "".join(f"<th>{district}</th>" for district in DISTRICTS)
        + f'</tr></thead><tbody>{table_rows}</tbody></table></div>'
    )
    return (
        '<!-- SHRINKAGE_DIAGNOSIS_START -->'
        '<section class="shrinkage-diagnosis" aria-labelledby="shrinkage-diagnosis-title">'
        '<div class="shrinkage-heading"><div><h2 id="shrinkage-diagnosis-title">2. 인구밀집구역(DID) 탈락지역 축소양상 진단</h2></div>'
        '<span class="shrinkage-score-badge"><i class="shrinkage-busan-key"></i>빨간 선: 부산 전체</span></div>'
        f'<div class="shrinkage-radar-grid">{"".join(rows)}</div>'
        + table +
        '<div class="shrinkage-footnote"><ul class="shrinkage-guide-list"><li>점수는 네 구와 부산 전체를 기준으로 지표별 표준화한 0–100 상대값이며, 공통 최댓값은 100.</li></ul></div>'
        '</section><!-- SHRINKAGE_DIAGNOSIS_END -->'
    )


STYLE = r'''/* SHRINKAGE_DIAGNOSIS_STYLE_START */
.shrinkage-diagnosis{margin:0 0 14px;padding:18px;background:#fff;border:1px solid #dfe7df;border-radius:14px;box-shadow:0 5px 20px rgba(36,61,48,.06)}
.shrinkage-heading{display:flex;justify-content:space-between;align-items:flex-start;gap:16px;margin-bottom:10px}
.shrinkage-heading h2{margin:0;color:#1d4136;font-size:clamp(20px,1.65vw,25px);letter-spacing:-.02em}
.shrinkage-heading p{margin:4px 0 0;color:#697d74;font-size:13px}
.shrinkage-score-badge{white-space:nowrap;border:1px solid #dce6dc;border-radius:999px;padding:6px 10px;color:#527064;font-size:11px}
.shrinkage-busan-key{display:inline-block;width:20px;margin-right:7px;vertical-align:middle;border-top:2px dashed #d7443e}
.shrinkage-radar-grid{display:grid;grid-template-columns:repeat(4,minmax(0,1fr));gap:12px}
.shrinkage-radar-card{min-width:0;border:1px solid #e1e8e1;border-top:3px solid var(--district-color);border-radius:11px;padding:10px 12px;background:#fff}
.shrinkage-radar-card header{display:flex;justify-content:space-between;align-items:baseline;gap:10px}
.shrinkage-radar-card h3{margin:0;color:var(--district-color);font-size:17px}
.shrinkage-radar-card header small{color:#7c8b82;font-size:10px}
.radar-svg{display:block;width:100%;height:auto;max-height:340px;overflow:visible}
.radar-grid{fill:none;stroke:#e2e9e3;stroke-width:1}
.radar-axis{stroke:#d8e1d9;stroke-width:1}
.radar-shape{fill:var(--radar-color);fill-opacity:.22;stroke:var(--radar-color);stroke-width:2.5}
.radar-busan-shape{fill:none;stroke:#d7443e;stroke-width:2.6;stroke-dasharray:7 4;stroke-linejoin:round}
.radar-busan-dot{fill:#d7443e;stroke:#fff;stroke-width:1.3}
.radar-dot{fill:var(--radar-color);stroke:#fff;stroke-width:1.5}
.radar-label{fill:#385348;font-size:19px;font-weight:750}
.shrinkage-metrics-table-wrap{margin-top:12px;overflow-x:auto;border:1px solid #e2eae2;border-radius:9px}
.shrinkage-metrics-table{width:100%;border-collapse:collapse;font-size:14px;line-height:1.45}
.shrinkage-metrics-table th,.shrinkage-metrics-table td{padding:8px 10px;border-bottom:1px solid #e8eee8;text-align:center;white-space:nowrap}
.shrinkage-metrics-table thead th{background:#f3f7f3;color:#355346;font-weight:700}
.shrinkage-metrics-table tbody th{text-align:left;color:#24483a}
.shrinkage-metrics-table .shrinkage-indicator{text-align:left;color:#63776c}
.shrinkage-metrics-table tbody tr:last-child th,.shrinkage-metrics-table tbody tr:last-child td{border-bottom:0}
@media(max-width:700px){.shrinkage-metrics-table{font-size:12px}.shrinkage-metrics-table th,.shrinkage-metrics-table td{padding:7px 8px}}
.shrinkage-footnote{margin:11px 1px 0;color:#738178;font-size:10.5px;line-height:1.6}
@media(max-width:1200px){.shrinkage-radar-grid{grid-template-columns:repeat(2,minmax(0,1fr))}}
@media(max-width:650px){.shrinkage-radar-grid{grid-template-columns:1fr}.shrinkage-heading{display:block}.shrinkage-score-badge{display:inline-block;margin-top:8px}.radar-svg{max-height:none}}
/* SHRINKAGE_DIAGNOSIS_STYLE_END */'''


def update_preview(data):
    html = PREVIEW.read_text(encoding="utf-8")
    section = html_section(data)
    start, end = "<!-- SHRINKAGE_DIAGNOSIS_START -->", "<!-- SHRINKAGE_DIAGNOSIS_END -->"
    if start in html and end in html:
        left = html.index(start)
        right = html.index(end, left) + len(end)
        html = html[:left] + section + html[right:]
    else:
        anchor = "<!-- DID_DASHBOARD_END -->"
        if anchor not in html:
            raise ValueError("Could not find DID dashboard end marker in preview")
        html = html.replace(anchor, anchor + "\n" + section, 1)
    style_start, style_end = "/* SHRINKAGE_DIAGNOSIS_STYLE_START */", "/* SHRINKAGE_DIAGNOSIS_STYLE_END */"
    if style_start in html and style_end in html:
        left = html.index(style_start)
        right = html.index(style_end, left) + len(style_end)
        html = html[:left] + STYLE + html[right:]
    else:
        html = html.replace("</style>", STYLE + "\n</style>", 1)
    PREVIEW.write_text(html, encoding="utf-8")


def main():
    DERIVED.mkdir(parents=True, exist_ok=True)
    data = make_data()
    (DERIVED / "shrinkage_diagnosis.json").write_text(
        json.dumps(data, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    update_preview(data)
    print(json.dumps({
        "districts": {
            district: {
                "population_growth": data["districts"][district]["population_growth"],
                "worker_growth": data["districts"][district]["worker_growth"],
                "living_ratio": data["districts"][district]["living_ratio"],
                "persistence": data["districts"][district]["persistence"],
                "connectivity": data["districts"][district]["connectivity"],
            }
            for district in DISTRICTS
        }
    }, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
