import collections
import json
import re
from pathlib import Path

import openpyxl

DATA = Path(__file__).resolve().parents[1] / "data"
OUT = DATA / "derived" / "busan_dong_function_change.json"


def normalize(value):
    return re.sub(r"\s+", "", value).replace("\uc81c", "")


employment = {"2021": {}, "2024": {}}
for year in employment:
    path = next(DATA.glob(f"*{year}\ub144\uae30\uc900\uc790\ub8cc.xlsx"))
    sheet = openpyxl.load_workbook(path, read_only=True, data_only=True).active
    current_gu = ""
    for row in sheet.iter_rows(min_row=5, values_only=True):
        region = str(row[0] or "").strip()
        industry = str(row[1] or "").strip()
        match = re.match(r"(\d+)\.(.*)", region)
        if not match:
            continue
        code, name = match.group(1), match.group(2).strip()
        if len(code) == 6:
            current_gu = name
        if len(code) in (7, 8) and industry.endswith("\uc0b0\uc5c5"):
            try:
                employment[year][normalize(current_gu + "|" + name)] = float(row[3])
            except (TypeError, ValueError):
                pass

living_path = next(DATA.glob("*\uc0dd\ud65c\uc778\uad6c*\uc6d4\ubcc4 \uc77c\ud3c9\uade0.xlsx"))
sheet = openpyxl.load_workbook(living_path, read_only=True, data_only=True).active
monthly = collections.defaultdict(lambda: [0.0, 0])
names = {}
for month, code, name, hour, home, work, visit in sheet.iter_rows(min_row=2, values_only=True):
    month, code = str(month), str(code)
    if month[:4] not in ("2023", "2024"):
        continue
    try:
        value = sum(float(item or 0) for item in (home, work, visit))
    except (TypeError, ValueError):
        continue
    key = (code, month)
    monthly[key][0] += value
    monthly[key][1] += 1
    names[code] = str(name).strip()

annual_months = collections.defaultdict(dict)
for (code, month), (total, count) in monthly.items():
    annual_months[code][month] = total / count

stats = {}
for code, full_name in names.items():
    bits = full_name.split(maxsplit=1)
    if len(bits) != 2:
        continue
    key = normalize(bits[0] + "|" + bits[1])
    values = {}
    for year in ("2023", "2024"):
        monthly_values = [value for month, value in annual_months[code].items() if month.startswith(year)]
        values[year] = sum(monthly_values) / len(monthly_values) if monthly_values else None
    pop23, pop24 = values["2023"], values["2024"]
    emp21, emp24 = employment["2021"].get(key), employment["2024"].get(key)
    stats[code] = {
        "name": full_name,
        "living2023": pop23,
        "living2024": pop24,
        "livingChangePct": (pop24 / pop23 - 1) * 100 if pop23 else None,
        "employment2021": emp21,
        "employment2024": emp24,
        "employmentChangePct": (emp24 / emp21 - 1) * 100 if emp21 else None,
    }

OUT.write_text(json.dumps(stats, ensure_ascii=False), encoding="utf-8")
matched_employment = sum(item["employmentChangePct"] is not None for item in stats.values())
print(f"Wrote {len(stats)} dongs; employment trend matched for {matched_employment}.")
